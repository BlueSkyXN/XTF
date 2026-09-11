"""收尾缺陷回归：保留真实配置、读取、规划和执行路径，仅替换远端 I/O。"""

from argparse import Namespace
from copy import deepcopy
from dataclasses import replace
import json
from unittest.mock import Mock

from openpyxl import Workbook
import pandas as pd
import pytest
import yaml

from api import A1Range, BaseV3Backend, BitableV1Backend
from api.base import RateLimiter, RetryableAPIClient
from core.config import SyncMode
from core.control import FixedWaitRetry, RequestController, RetryConfig
from core.plan import OutcomeStatus
from tests.conftest import attach_runtime
from tests.test_plan import make_file_bitable_engine
from tests.test_review_regressions import sheet_service
from xtf_cli.config import make_template, resolve_config
from xtf_cli.runtime import _load_dataframe, main


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    monkeypatch.setattr(
        "requests.request", Mock(side_effect=AssertionError("network forbidden"))
    )


@pytest.mark.parametrize("existing", [False, True])
def test_invalid_init_never_truncates_or_creates_config(tmp_path, capsys, existing):
    path = tmp_path / "config.yaml"
    original = "sentinel: keep-existing-config\n"
    if existing:
        path.write_text(original, encoding="utf-8")
    code = main(
        [
            "config",
            "init",
            "--source-type",
            "bitable",
            "--target-type",
            "sheet",
            "--output",
            str(path),
            "--force",
            "--json",
        ]
    )
    assert code == 3
    assert json.loads(capsys.readouterr().out)["ok"] is False
    assert (
        path.read_text(encoding="utf-8") == original if existing else not path.exists()
    )


def test_init_replace_failure_preserves_config_and_removes_temporary(
    tmp_path, monkeypatch, capsys
):
    path = tmp_path / "config.yaml"
    path.write_text("original: content\n", encoding="utf-8")
    monkeypatch.setattr(
        "xtf_cli.config.os.replace", Mock(side_effect=OSError("write denied"))
    )
    assert main(["config", "init", "--output", str(path), "--force", "--json"]) == 3
    capsys.readouterr()
    assert path.read_text(encoding="utf-8") == "original: content\n"
    assert list(tmp_path.iterdir()) == [path]


@pytest.mark.parametrize("selector,expected", [("0", "named"), (0, "first")])
def test_yaml_sheet_selector_preserves_name_and_index(tmp_path, selector, expected):
    source = tmp_path / "tabs.xlsx"
    book = Workbook()
    book.active.title = "first"
    book.active.append(["ID"])
    book.active.append(["first"])
    sheet = book.create_sheet("0")
    sheet.append(["ID"])
    sheet.append(["named"])
    book.save(source)
    document = make_template("file", "bitable")
    document["auth"]["app_secret"] = "fake-secret"
    document["source"]["file"] = {"path": str(source), "sheet_name": selector}
    path = tmp_path / "config.yaml"
    path.write_text(yaml.safe_dump(document), encoding="utf-8")
    resolved = resolve_config(Namespace(config=str(path)), environ={})
    assert _load_dataframe(resolved)["ID"].tolist() == [expected]


@pytest.mark.parametrize("failure", ["corrupt", "missing_sheet"])
def test_excel_read_failure_keeps_input_exit_code(tmp_path, capsys, failure):
    source = tmp_path / "source.xlsx"
    if failure == "corrupt":
        source.write_bytes(b"not-an-xlsx-zip")
    else:
        book = Workbook()
        book.active.append(["ID"])
        book.save(source)
    document = make_template("file", "bitable")
    document["auth"]["app_secret"] = "fake-secret"
    document["source"]["file"] = {
        "path": str(source),
        "sheet_name": "absent" if failure == "missing_sheet" else None,
    }
    path = tmp_path / "config.yaml"
    path.write_text(yaml.safe_dump(document), encoding="utf-8")
    assert main(["sync", "--config", str(path), "--dry-run", "--json"]) == 3
    assert json.loads(capsys.readouterr().out)["error"]["code"] == "XTF_E_INPUT_READ"


@pytest.mark.parametrize("status", [429, 503])
@pytest.mark.parametrize("max_wait", [None, 2.0])
def test_advanced_retry_respects_server_delay_and_total_budget(
    monkeypatch, status, max_wait
):
    limited = Mock(status_code=status, headers={"Retry-After": "5"})
    ok = Mock(status_code=200, headers={})
    request = Mock(side_effect=[limited, ok])
    sleep = Mock()
    monkeypatch.setattr("requests.request", request)
    monkeypatch.setattr("time.sleep", sleep)
    controller = RequestController(
        FixedWaitRetry(
            RetryConfig(initial_delay=0.5, max_retries=1, max_wait_time=max_wait)
        )
    )
    client = RetryableAPIClient(
        max_retries=1, rate_limiter=RateLimiter(0), controller=controller
    )
    response = client.call_api("GET", "https://example.invalid/test")
    if max_wait is None:
        assert response is ok
        sleep.assert_called_once_with(5.0)
        assert request.call_count == 2
    else:
        assert response is limited
        sleep.assert_not_called()
        assert request.call_count == 1


@pytest.mark.parametrize("empty", [False, True])
@pytest.mark.parametrize("offset", [False, True])
def test_sheet_clone_uses_real_full_range_and_clears_outside_start(empty, offset):
    state = [[""] * 8 for _ in range(20)]
    sr, sc = (3, 2) if offset else (1, 1)
    state[0][0] = "outside-sentinel"
    state[sr - 1][sc - 1 : sc + 1] = ["ID", "Name"]
    state[sr][sc - 1 : sc + 1] = [1, "old"]
    engine = sheet_service(
        [],
        mode="clone",
        start_row=sr,
        start_column="B" if offset else "A",
        verify_remote_writes=True,
        verify_timeout_seconds=0,
    )

    def read(token, text, **kwargs):
        area = A1Range.parse(text)
        return deepcopy(
            [
                row[area.start_col - 1 : area.end_col]
                for row in state[area.start_row - 1 : area.end_row]
            ]
        )

    def chunked(token, sid, start_row, end_row, start_col, end_col, **kwargs):
        return read(token, f"{sid}!{start_col}{start_row}:{end_col}{end_row}")

    def mutate(method, token, endpoint, body, **kwargs):
        responses = []
        for item in body.get("valueRanges", [body.get("valueRange")]):
            area = A1Range.parse(item["range"])
            for row, values in enumerate(item["values"], area.start_row - 1):
                state[row][area.start_col - 1 : area.end_col] = deepcopy(values)
            responses.append({"updatedRange": area.text})
        return {"code": 0, "data": {"responses": responses}}

    engine.api.get_sheet_grid_properties = Mock(return_value=(20, 8))
    engine.api.get_sheet_data = Mock(side_effect=read)
    engine.api.get_sheet_data_chunked = Mock(side_effect=chunked)
    engine.api._typed_values_call = Mock(side_effect=mutate)
    engine._setup_sheet_intelligence = Mock(return_value=True)
    frame = pd.DataFrame() if empty else pd.DataFrame({"ID": [2], "Name": ["new"]})
    plan = engine.plan(frame)
    assert plan.actions[0].a1_range == f"{engine.runtime.target.sheet_id}!A1:H20"
    result = engine.execute_plan(plan)
    assert result.ok, result.error
    if empty:
        assert all(not cell for row in state for cell in row)
    else:
        assert state[sr - 1][sc - 1 : sc + 1] == ["ID", "Name"]
        assert state[sr][sc - 1 : sc + 1] == [2, "new"]
    if offset:
        assert state[0][0] == ""


@pytest.mark.parametrize("revision_changes", [False, True])
@pytest.mark.parametrize("interference", [None, "record", "schema"])
@pytest.mark.parametrize("visibility_delay", [False, True])
@pytest.mark.parametrize("backend_kind", ["base_v3", "bitable_v1"])
def test_bitable_clone_advances_only_its_own_new_fields(
    revision_changes, interference, visibility_delay, backend_kind, monkeypatch
):
    engine = make_file_bitable_engine(
        SyncMode.CLONE,
        create_missing_fields=True,
        index_column=None,
        backend=backend_kind,
    )
    attach_runtime(
        engine,
        replace(
            engine.runtime,
            sync=replace(
                engine.sync_config, verify_timeout_seconds=1 if visibility_delay else 0
            ),
        ),
    )
    monkeypatch.setattr("core.verification.sleep", lambda _: None)
    v1 = backend_kind == "bitable_v1"
    backend = BitableV1Backend(Mock()) if v1 else BaseV3Backend(Mock())
    fields = [{"id": "fld_id", "name": "ID", "type": "number"}]
    records = {"rec_old": {"ID": 1}}
    writes = []
    revision = 100
    delayed = False

    def response(data):
        return {"code": 0, "data": data} if v1 else data

    def call(method, url, **kwargs):
        nonlocal revision, delayed
        if method == "GET" and url.endswith("/fields"):
            if v1:
                return response(
                    {
                        "items": [
                            {
                                "field_id": f["id"],
                                "field_name": f["name"],
                                "type": 2 if f["type"] == "number" else 1,
                                "property": (
                                    {"description": f["description"]}
                                    if "description" in f
                                    else {}
                                ),
                            }
                            for f in fields
                        ],
                        "has_more": False,
                    }
                )
            return {"fields": list(fields), "total": len(fields)}
        if method == "POST" and url.endswith("/fields"):
            body = kwargs["json"]
            field = {
                "id": f"fld_{len(fields)}",
                **(
                    {
                        "name": body["field_name"],
                        "type": "number" if body["type"] == 2 else "text",
                    }
                    if v1
                    else body
                ),
            }
            fields.append(field)
            writes.append("field")
            if revision_changes:
                revision += 1
            if interference == "record":
                records["rec_old"]["ID"] = 999
            elif interference == "schema":
                fields[0] = {**fields[0], "description": "external change"}
            return response({"field": field})
        if (method == "GET" and url.endswith("/records")) or (
            v1 and url.endswith("/records/search")
        ):
            visible_fields = fields
            if visibility_delay and writes == ["field"] and not delayed:
                visible_fields = fields[:1]
                delayed = True
            if v1:
                assert "field_names" not in kwargs["json"]
                return response(
                    {
                        "items": [
                            {
                                "record_id": rid,
                                "fields": {
                                    f["name"]: record.get(f["name"])
                                    for f in visible_fields
                                },
                            }
                            for rid, record in records.items()
                        ],
                        "has_more": False,
                    }
                )
            return {
                "fields": [f["name"] for f in visible_fields],
                "field_id_list": [f["id"] for f in visible_fields],
                "field_type_list": [f["type"] for f in visible_fields],
                "data": [
                    [r.get(f["name"]) for f in visible_fields] for r in records.values()
                ],
                "record_id_list": list(records),
                "has_more": False,
                "timezone": "UTC",
                "rev": revision,
            }
        if method == "POST" and url.endswith("/batch_delete"):
            ids = kwargs["json"]["records" if v1 else "record_id_list"]
            for rid in ids:
                records.pop(rid)
            writes.append("delete")
            return response(
                {
                    "revision": revision,
                    "records": [{"record_id": rid, "deleted": True} for rid in ids],
                }
            )
        if method == "POST" and url.endswith("/batch_create"):
            new_ids = []
            rows = (
                [r["fields"] for r in kwargs["json"]["records"]]
                if v1
                else kwargs["json"]["create_records"]
            )
            for row in rows:
                rid = f"rec_new_{len(records)}"
                records[rid] = row
                new_ids.append(rid)
            writes.append("create")
            return response({"record_id_list": new_ids, "revision": revision})
        raise AssertionError((method, url))

    backend._call = Mock(side_effect=call)
    engine.api = backend
    result = engine.execute_plan(
        engine.plan(pd.DataFrame({"ID": [2], "Name": ["new"], "Score": [3]}))
    )
    if interference is None:
        assert result.ok, result.error
        assert writes == ["field", "field", "delete", "create"]
        assert list(records.values()) == [{"ID": 2, "Name": "new", "Score": 3}]
    else:
        assert result.status is OutcomeStatus.PARTIAL
        assert writes == ["field"]
        assert "rec_old" in records


@pytest.mark.parametrize("invalid_id", ["", "   ", None])
def test_invalid_created_id_stops_next_append_only_batch(invalid_id):
    engine = make_file_bitable_engine(
        SyncMode.INCREMENTAL,
        match_strategy="append_only",
        index_column=None,
        batch_size=1,
    )
    backend = BaseV3Backend(Mock())
    sent = []

    def call(method, url, **kwargs):
        if method == "GET" and url.endswith("/fields"):
            return {
                "fields": [{"id": "fld_name", "name": "Name", "type": "text"}],
                "total": 1,
            }
        if method == "POST" and url.endswith("/batch_create"):
            sent.append(kwargs["json"])
            return {"record_id_list": [invalid_id if len(sent) == 1 else "rec_next"]}
        raise AssertionError((method, url))

    backend._call = Mock(side_effect=call)
    engine.api = backend
    result = engine.execute_plan(
        engine.plan(pd.DataFrame({"Name": ["first", "second"]}))
    )
    assert result.status is OutcomeStatus.INDETERMINATE
    assert len(sent) == 1

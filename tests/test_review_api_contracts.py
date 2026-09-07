"""Adversarial wire-contract and whole-execution regressions (2026-09-05)."""

from dataclasses import replace
from unittest.mock import Mock, patch

import pandas as pd
import pytest

from api import A1Range, FeishuAPIError, SheetAPI
from api.base import RateLimiter, RetryableAPIClient
from api.bitable_backend import (
    BitableBackendKind,
    CanonicalRecord,
    MutationOutcome,
    MutationReceipt,
    RecordReadResult,
)
from api.bitable_v1 import BitableV1Backend
from api.bitable_v3 import BaseV3Backend
from api.sdk import FeishuResponseParser
from core.config import SyncMode
from core.plan import (
    AppendRowsAction,
    OutcomeStatus,
    WriteColumnsAction,
    WriteRangeAction,
)
from core.snapshot import SourceTable
from tests.test_plan import make_file_bitable_engine
from tests.test_review_regressions import sheet_service


def response(data=None, *, code=0, status=200, headers=None):
    value = Mock()
    value.status_code = status
    value.headers = headers or {}
    value.json.return_value = {"code": code, "data": {} if data is None else data}
    return value


@pytest.mark.parametrize("backend_type", [BitableV1Backend, BaseV3Backend])
@pytest.mark.parametrize("advanced", [False, True])
def test_review_mixed_http_business_retries_share_one_total_budget(
    backend_type, advanced
):
    controller = None
    if advanced:
        from core.control import build_request_controller

        controller = build_request_controller(
            retry_config={"initial_delay": 0, "max_retries": 2},
            rate_limit_config={"delay": 0},
        )
    transport = RetryableAPIClient(
        max_retries=2,
        rate_limiter=RateLimiter(0),
        jitter_ratio=0,
        controller=controller,
    )
    backend = backend_type(Mock(), transport)
    # Three actual sends must exhaust the budget, not grant the business retry
    # another three transport attempts.
    replies = [
        response(status=503),
        response(code=1254290),
        response(status=503),
        response(),
    ]
    with patch("api.base.requests.request", side_effect=replies) as send, patch(
        "time.sleep"
    ):
        with pytest.raises(FeishuAPIError):
            backend._call("GET", "https://example.invalid/records")
    assert send.call_count == 3


def test_review_v3_business_error_preserves_retry_after_status_and_log_id():
    transport = Mock(max_retries=1)
    transport.call_api.side_effect = [
        response(code=1254290, headers={"Retry-After": "7", "X-Tt-Logid": "log-x"}),
        response(),
    ]
    with patch("time.sleep") as sleep:
        assert (
            BaseV3Backend(Mock(), transport)._call(
                "GET", "https://example.invalid/records"
            )
            == {}
        )
    sleep.assert_called_once_with(7.0)
    transport = Mock(max_retries=0)
    transport.call_api.return_value = response(
        code=1254290, headers={"X-Tt-Logid": "log-x"}
    )
    with pytest.raises(FeishuAPIError) as caught:
        BaseV3Backend(Mock(), transport)._call("GET", "https://example.invalid/records")
    assert caught.value.http_status == 200
    assert caught.value.log_id == "log-x"
    assert caught.value.retryable


def test_review_v3_create_field_does_not_replay_lost_response():
    transport = Mock(max_retries=2)
    transport.call_api.side_effect = FeishuAPIError.from_transport("response lost")
    receipt = BaseV3Backend(Mock(), transport).create_field("app", "table", "Name", 1)
    assert receipt.outcome is MutationOutcome.UNKNOWN_OUTCOME
    assert transport.call_api.call_count == 1
    assert transport.call_api.call_args.kwargs["retry_transport"] is False


@pytest.mark.parametrize("operation", ["batch_create", "batch_update"])
def test_review_v1_documented_records_response_supplies_confirmed_record_ids(operation):
    transport = Mock()
    transport.call_api.return_value = response(
        {"records": [{"record_id": "rec1", "fields": {"ID": 1}}]}
    )
    body = {"records": [{"record_id": "rec1", "fields": {"ID": 1}}]}
    receipt = BitableV1Backend(Mock(), transport)._mutation_call(
        operation, "app", "table", body, 1
    )
    assert receipt.outcome is MutationOutcome.ACCEPTED
    assert receipt.record_ids == ("rec1",)
    assert receipt.accepted_count == 1


def test_review_v1_deleted_false_is_not_counted_as_success():
    transport = Mock()
    transport.call_api.return_value = response(
        {
            "records": [
                {"record_id": "rec1", "deleted": True},
                {"record_id": "rec2", "deleted": False},
            ]
        }
    )
    receipt = BitableV1Backend(Mock(), transport).batch_delete(
        "app", "table", ["rec1", "rec2"]
    )
    assert receipt.outcome is MutationOutcome.PARTIAL
    assert receipt.accepted_count == 1
    assert receipt.record_ids == ("rec1",)
    assert receipt.record_not_found == ()
    assert receipt.raw_metadata["failed_record_ids"] == ("rec2",)


@pytest.mark.parametrize("backend_type", [BitableV1Backend, BaseV3Backend])
@pytest.mark.parametrize("failure", ["http_500", "invalid_json", "malformed_receipt"])
def test_review_sent_mutation_with_unusable_response_is_unknown(backend_type, failure):
    reply = (
        response(status=500)
        if failure == "http_500"
        else response({"records": "invalid", "record_id_list": 7})
    )
    if failure == "invalid_json":
        reply.json.side_effect = ValueError("truncated")
    transport = Mock(max_retries=0)
    transport.call_api.return_value = reply
    backend = backend_type(Mock(), transport)
    receipt = backend._mutation_call("batch_create", "app", "table", {"records": []}, 1)
    assert receipt.outcome is MutationOutcome.UNKNOWN_OUTCOME
    assert receipt.accepted_count == 0


@pytest.mark.parametrize(
    "payload", [{}, {"data": {}}, {"code": False}, {"code": None}, {"code": 0.1}]
)
def test_review_missing_or_invalid_success_code_is_not_success(payload):
    reply = response()
    reply.json.return_value = payload
    with pytest.raises(FeishuAPIError):
        FeishuResponseParser.parse(reply)


@pytest.mark.parametrize(
    "data",
    [
        {},
        {"valueRange": {}},
        {"valueRange": {"values": "bad"}},
        {"valueRange": {"values": [[1, 2, 3]]}},
    ],
)
def test_review_malformed_sheet_read_cannot_be_mistaken_for_empty(data):
    transport = Mock()
    transport.call_api.return_value = response(data)
    with pytest.raises(FeishuAPIError):
        SheetAPI(Mock(), transport).get_sheet_data("token", "sheet!A1:B2")


@pytest.mark.parametrize("status", [500, 502, 503])
def test_review_sheet_server_failure_may_have_applied(status):
    transport = Mock()
    transport.call_api.return_value = response(status=status)
    receipt = SheetAPI(Mock(), transport).append_values(
        "token", "sheet!A5:B5", [[1, "new"]]
    )
    assert receipt.outcome is MutationOutcome.UNKNOWN_OUTCOME
    assert receipt.raw_metadata["http_status"] == status


def test_review_pre_send_rate_failure_does_not_claim_unknown_remote_write():
    error = FeishuAPIError.from_transport("rate wait failed before send")
    error.response_data = {"request_started": False}
    transport = Mock()
    transport.call_api.side_effect = error
    receipt = SheetAPI(Mock(), transport).append_values(
        "token", "sheet!A5:B5", [[1, "new"]]
    )
    assert receipt.outcome is MutationOutcome.REJECTED


def test_review_sheet_auth_failure_keeps_auth_category():
    engine = sheet_service([])
    engine._reset_action_execution_state()
    receipt = SheetAPI._typed_failure_receipt(
        "write",
        1,
        0,
        [],
        FeishuAPIError(10003, "denied", http_status=403),
        failed_batch_index=1,
        raw_responses=[],
    )
    assert not engine._finalize_sheet_mutation(receipt)
    assert engine._last_action_error_kind.value == "auth"


def test_review_append_anchor_is_after_last_occupied_physical_row():
    engine = sheet_service(
        [["ID", "Name"], [1, "old"], [], [2, "last"]], mode="incremental"
    )
    plan = engine.plan(pd.DataFrame({"ID": [3, 4], "Name": ["new3", "new4"]}))
    append = next(
        action for action in plan.actions if isinstance(action, AppendRowsAction)
    )
    assert append.start_row == 5
    engine.api.append_values = Mock(
        return_value=MutationReceipt(
            "append",
            "sheet",
            2,
            accepted_count=2,
            unit="row",
            actual_ranges=(A1Range.parse("sheet!A5:B6"),),
        )
    )
    engine._reset_action_execution_state()
    assert engine._execute_action(append)
    assert engine.api.append_values.call_args.args[1].endswith("!A5:B6")


@pytest.mark.parametrize("with_index", [False, True])
def test_review_whole_clone_stops_on_concurrent_insert_during_post_delete_read(
    with_index,
):
    engine = make_file_bitable_engine(
        SyncMode.CLONE, backend="bitable_v1", index_column="ID" if with_index else None
    )
    fields = engine.api.list_fields.return_value
    original = RecordReadResult(
        (CanonicalRecord("old", {"ID": 1}),),
        fields,
        True,
        BitableBackendKind.BITABLE_V1,
    )
    concurrent = replace(original, records=(CanonicalRecord("concurrent", {"ID": 99}),))
    engine.api.list_records.return_value = original
    plan = engine.plan(pd.DataFrame({"ID": [2], "Name": ["new"]}))
    engine.api.list_records.side_effect = [original, concurrent]
    engine.api.batch_delete.return_value = MutationReceipt(
        "delete", BitableBackendKind.BITABLE_V1, 1, accepted_count=1, unit="record"
    )
    outcome = engine.execute_plan(plan)
    assert outcome.status is OutcomeStatus.PARTIAL
    engine.api.batch_create.assert_not_called()


def test_review_two_row_chunks_are_one_complete_column_not_two_columns():
    engine = sheet_service([])
    engine._reset_action_execution_state()
    action = WriteColumnsAction(
        column_data={"Name": ("a", "b", "c", "d"), "Other": (1, 2, 3, 4)},
        column_positions={"Name": 2, "Other": 3},
        start_row=2,
        max_gap=0,
        header_width=3,
        scope={},
    )
    engine._record_action_receipt(
        MutationReceipt(
            "write",
            "sheet",
            4,
            accepted_count=2,
            unit="range",
            actual_ranges=(A1Range.parse("sheet!B2:B3"), A1Range.parse("sheet!B4:B5")),
            outcome=MutationOutcome.PARTIAL,
        )
    )
    prefix = engine._applied_action_prefix(action)
    assert prefix.count == 1
    assert prefix.unit.value == "column"
    assert prefix.scope["accepted_units"] == 2
    assert prefix.scope["receipt_unit"] == "range"


def test_review_partial_column_chunk_does_not_claim_a_complete_column():
    engine = sheet_service([])
    engine._reset_action_execution_state()
    action = WriteColumnsAction(
        column_data={"Name": ("a", "b", "c", "d")},
        column_positions={"Name": 2},
        start_row=2,
        max_gap=0,
        header_width=2,
        scope={},
    )
    engine._record_action_receipt(
        MutationReceipt(
            "write",
            "sheet",
            2,
            accepted_count=1,
            unit="range",
            actual_ranges=(A1Range.parse("sheet!B2:B3"),),
            outcome=MutationOutcome.PARTIAL,
        )
    )
    prefix = engine._applied_action_prefix(action)
    assert prefix.count == 0
    assert prefix.scope["applied_physical_rows"] == 2


def test_review_source_large_integer_is_not_upcast_by_neighbor_float_column():
    frame = SourceTable.from_dataframe(
        pd.DataFrame({"ID": [2**60 + 1], "Amount": [1.25]})
    ).to_dataframe()
    assert next(frame.iterrows())[1]["ID"] == 2**60 + 1


def test_review_append_inserts_rows_and_anchors_next_chunk_after_actual_result():
    transport = Mock()
    transport.call_api.side_effect = [
        response({"updates": {"updatedRange": "sheet!A10:A10"}}),
        response({"updates": {"updatedRange": "sheet!A11:A11"}}),
    ]
    api = SheetAPI(Mock(), transport, write_max_rows=1)
    receipt = api.append_values("token", "sheet!A5:A6", [[1], [2]])
    assert receipt.outcome is MutationOutcome.ACCEPTED
    assert transport.call_api.call_args_list[0].kwargs["params"] == {
        "insertDataOption": "INSERT_ROWS"
    }
    assert (
        transport.call_api.call_args_list[1].kwargs["json"]["valueRange"]["range"]
        == "sheet!A11:A11"
    )


def test_review_overwrite_preserves_blank_physical_columns_and_unmatched_cells():
    engine = sheet_service(
        [["Name", None, "ID"], ["keep", "scratch", 1], [], ["old", "remove", 2]],
        mode="overwrite",
        start_row=3,
        start_column="B",
    )
    plan = engine.plan(pd.DataFrame({"ID": [2], "Name": ["new"]}))
    write = next(
        action for action in plan.actions if isinstance(action, WriteRangeAction)
    )
    assert write.values == (
        ("Name", None, "ID"),
        ("keep", "scratch", 1),
        ("", "", ""),
        ("new", "", 2),
    )


@pytest.mark.parametrize("suffix", ["xlsx", "csv"])
def test_review_real_file_preserves_text_ids_and_na_literals(tmp_path, suffix):
    from core.reader import DataFileReader

    path = tmp_path / f"data.{suffix}"
    if suffix == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        for row in (("ID", "Name"), ("001", "NA"), ("002", "NULL")):
            workbook.active.append(row)
        workbook.save(path)
    else:
        path.write_text("ID,Name\n001,NA\n002,NULL\n", encoding="utf-8")
    data = DataFileReader().read_file(path)
    assert data.to_dict(orient="list") == {"ID": ["001", "002"], "Name": ["NA", "NULL"]}


@pytest.mark.parametrize("header", [("ID", "ID"), ("ID", "")])
def test_review_real_excel_does_not_mangle_ambiguous_headers(tmp_path, header):
    from openpyxl import Workbook
    from core.reader import DataFileReader

    path = tmp_path / "bad.xlsx"
    workbook = Workbook()
    workbook.active.append(header)
    workbook.active.append((1, "value"))
    workbook.save(path)
    with pytest.raises(ValueError):
        DataFileReader().read_file(path)


def test_review_duplicate_yaml_keys_are_rejected(tmp_path):
    from xtf_cli.config import read_v2_file
    from xtf_cli.errors import CLIError

    path = tmp_path / "xtf.yaml"
    path.write_text(
        "schema_version: 2\nsync:\n  mode: full\n  mode: clone\n", encoding="utf-8"
    )
    with pytest.raises(CLIError, match="duplicate configuration key"):
        read_v2_file(path)


@pytest.mark.parametrize("field", ["app_id", "start_column"])
def test_review_null_config_does_not_become_literal_none(field):
    from core.config import TargetType
    from tests.conftest import make_runtime_config

    with pytest.raises(ValueError):
        make_runtime_config(TargetType.SHEET, index_column="ID", **{field: None})


@pytest.mark.parametrize(
    "field",
    [
        "rate_limit_delay",
        "retry_initial_delay",
        "retry_max_wait_time",
        "retry_multiplier",
        "retry_increment",
        "rate_limit_window_size",
    ],
)
@pytest.mark.parametrize("value", [float("nan"), float("inf"), float("-inf")])
def test_review_nonfinite_retry_and_rate_values_rejected(field, value):
    from core.config import TargetType
    from tests.conftest import make_runtime_config

    with pytest.raises(ValueError, match="有限"):
        make_runtime_config(TargetType.SHEET, index_column="ID", **{field: value})


def test_review_cli_real_excel_real_planner_dry_run_never_executes(
    tmp_path, monkeypatch, capsys
):
    import json
    from openpyxl import Workbook
    from xtf_cli import main

    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    workbook.active.append(("ID", "Name"))
    workbook.active.append((2, "new"))
    workbook.save(path)
    engine = sheet_service([["Name", "ID"], ["old", 1]], mode="incremental")
    engine.execute_plan = Mock(side_effect=AssertionError("dry-run must never execute"))
    monkeypatch.setattr("core.service.SyncService", lambda config: engine)
    with patch("requests.request", side_effect=AssertionError("unexpected live HTTP")):
        code = main(
            [
                "sync",
                "--file",
                str(path),
                "--source-type",
                "file",
                "--target-type",
                "sheet",
                "--app-id",
                "app",
                "--app-secret",
                "secret",
                "--spreadsheet-token",
                "token",
                "--sheet-id",
                "sheet",
                "--mode",
                "incremental",
                "--match-strategy",
                "by_key",
                "--index-column",
                "ID",
                "--dry-run",
                "--json",
            ]
        )
    captured = capsys.readouterr()
    assert code == 0, captured.err or captured.out
    result = json.loads(captured.out)
    assert result["dry_run"] is True
    assert any(action["kind"] == "append_rows" for action in result["plan"]["actions"])
    engine.execute_plan.assert_not_called()


@pytest.mark.parametrize("data", [{}, {"items": []}, {"has_more": False}])
def test_review_v1_incomplete_page_envelope_cannot_mean_empty_source(data):
    with pytest.raises(FeishuAPIError):
        BitableV1Backend._page_data({"code": 0, "data": data})


@pytest.mark.parametrize(
    "token,expire",
    [
        (None, 7200),
        ("", 7200),
        (42, 7200),
        ("token", float("nan")),
        ("token", -1),
        ("token", True),
    ],
)
def test_review_auth_rejects_malformed_token_response(token, expire):
    from api.auth import FeishuAuth

    transport = Mock()
    reply = response()
    reply.json.return_value = {
        "code": 0,
        "tenant_access_token": token,
        "expire": expire,
    }
    transport.call_api.return_value = reply
    auth = FeishuAuth("app", "secret", transport)
    with pytest.raises(FeishuAPIError):
        auth.get_auth_headers()
    assert auth.tenant_access_token is None


def test_review_auth_http_failure_not_overridden_by_success_code():
    from api.auth import FeishuAuth

    reply = response(status=503)
    reply.json.return_value = {
        "code": 0,
        "tenant_access_token": "token",
        "expire": 7200,
    }
    transport = Mock()
    transport.call_api.return_value = reply
    with pytest.raises(FeishuAPIError) as caught:
        FeishuAuth("app", "secret", transport).get_auth_headers()
    assert caught.value.http_status == 503


def test_review_auth_reuses_valid_cached_token():
    from api.auth import FeishuAuth

    reply = response()
    reply.json.return_value = {
        "code": 0,
        "tenant_access_token": "token",
        "expire": 7200,
    }
    transport = Mock()
    transport.call_api.return_value = reply
    auth = FeishuAuth("app", "secret", transport)
    assert auth.get_auth_headers()["Authorization"] == "Bearer token"
    assert auth.get_auth_headers()["Authorization"] == "Bearer token"
    transport.call_api.assert_called_once()


@pytest.mark.parametrize(
    "backend_type,data", [(BitableV1Backend, {"records": []}), (BaseV3Backend, {})]
)
def test_review_empty_create_acknowledgement_does_not_claim_nothing_was_written(
    backend_type, data
):
    transport = Mock(max_retries=0)
    transport.call_api.return_value = response(data)
    receipt = backend_type(Mock(), transport)._mutation_call(
        "batch_create", "app", "table", {"records": []}, 1
    )
    assert receipt.outcome is MutationOutcome.UNKNOWN_OUTCOME


def test_review_post_write_read_does_not_adopt_concurrent_edit_to_next_target_row():
    values = [["ID", "Name"], [1, "old1"], [], [2, "old2"]]
    engine = sheet_service(values)
    plan = engine.plan(pd.DataFrame({"ID": [1, 2], "Name": ["new1", "new2"]}))
    calls = []

    def execute(action):
        calls.append(action)
        values[1] = [1, "new1"]
        values[3] = [2, "concurrent"]
        engine._record_action_receipt(
            MutationReceipt(
                "write",
                "sheet",
                1,
                accepted_count=1,
                unit="range",
                actual_ranges=(A1Range.parse("sheet!A2:B2"),),
            )
        )
        engine._last_action_mutation_complete = True
        return True

    engine._execute_action = execute
    outcome = engine.execute_plan(plan)
    assert outcome.status is OutcomeStatus.PARTIAL
    assert len(calls) == 1
    assert values[3] == [2, "concurrent"]


def test_review_formula_verification_width_includes_unnamed_physical_gaps():
    engine = sheet_service(
        [["ID", None, "Name"], [1, "untouched", "old"]], start_column="B"
    )
    plan = engine.plan(pd.DataFrame({"ID": [1], "Name": ["new"]}))
    patch_action = next(
        action for action in plan.actions if isinstance(action, WriteColumnsAction)
    )
    assert patch_action.header_width == 3


@pytest.mark.parametrize(
    "literal,expected", [("1e3", 1000), ("1E-3", 0.001), ("2.5e2", 250)]
)
def test_review_non_key_scientific_number_does_not_fall_back_to_mantissa(
    literal, expected
):
    engine = sheet_service([])
    assert engine.converter._force_to_number(literal, "Amount") == expected
